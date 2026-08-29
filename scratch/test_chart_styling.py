import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.axis import NumericAxis

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Test"

# Dummy data
ws.append(["Seq", "MQ", "Limit"])
for i in range(1, 27):
    ws.append([i, 0.5 + 0.3 * (i % 3), 2.0])

chart = ScatterChart()
chart.title = "Test MQ Chart"
chart.style = 13
chart.width = 22
chart.height = 11.5

xvalues = Reference(ws, min_col=1, min_row=2, max_row=27)
yvalues_mq = Reference(ws, min_col=2, min_row=2, max_row=27)
s_mq = Series(yvalues_mq, xvalues, title="MQ Blank (n=26)")
s_mq.marker = Marker('circle')
s_mq.marker.size = 6
s_mq.marker.graphicalProperties.solidFill = "0284C7"
s_mq.marker.graphicalProperties.line.solidFill = "0369A1"
s_mq.graphicalProperties.line.noFill = True
chart.series.append(s_mq)

yvalues_lim = Reference(ws, min_col=3, min_row=2, max_row=27)
s_lim = Series(yvalues_lim, xvalues, title="2.0 uM Limit")
s_lim.marker.symbol = "none"
s_lim.graphicalProperties.line.solidFill = "DC2626"
s_lim.graphicalProperties.line.width = 20000  # in EMUs (1 pt = 12700 EMUs)
s_lim.graphicalProperties.line.prstDash = "dash"
chart.series.append(s_lim)

if chart.legend:
    chart.legend.legendPos = "tr"

chart.x_axis.title = "Sequence (1-26)"
chart.y_axis.title = "DOC (umol/L)"
chart.y_axis.scaling.min = 0.0
chart.y_axis.scaling.max = 3.5

ws.add_chart(chart, "E2")
wb.save("scratch/test_styled_chart.xlsx")
print("Successfully generated test_styled_chart.xlsx")
