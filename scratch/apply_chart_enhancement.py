import os
import sys
import shutil
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties

sys.stdout.reconfigure(encoding='utf-8')

target_file = r'F:\印度洋测样\ODV\202608\20260827\Ocean_DOC_MultiColumn_QC_Report_GEOMAR_Validated_v2_Processed_latest.xlsx'
if not os.path.exists(target_file):
    target_file = r'F:\印度洋测样\ODV\202608\20260826\Ocean_DOC_MultiColumn_QC_Report_GEOMAR_Validated_v2_Processed_latest.xlsx'

print(f"Target file: {target_file}")

# 1. Create a safe backup before any modification
backup_path = target_file.replace('.xlsx', '_backup_before_chart_upgrade.xlsx')
shutil.copy2(target_file, backup_path)
print(f"Safety backup created at: {backup_path}")

# 2. Load workbook
wb = openpyxl.load_workbook(target_file, data_only=False)
ws = wb['Executive_Dashboard']

print(f"Original charts count: {len(ws._charts)}")
# Clear existing charts on Executive_Dashboard
ws._charts.clear()

# -----------------------------------------------------------------------------
# Chart 1: MQ Baseline Scatter Chart
# -----------------------------------------------------------------------------
chart_mq = ScatterChart()
chart_mq.title = "【图 1】全航段 26 序列 / 16 柱 MQ 空白基线总览散点图 (Live MQ Baseline)"
chart_mq.style = 13
chart_mq.width = 22.0
chart_mq.height = 11.5

# MQ Data points (n=513)
x_mq = Reference(ws, min_col=1, min_row=62, max_row=574)
y_mq = Reference(ws, min_col=3, min_row=62, max_row=574)
s_mq = Series(y_mq, x_mq, title="MQ 空白测定点 (n=513)")
s_mq.marker = Marker('circle')
s_mq.marker.size = 5
s_mq.marker.graphicalProperties.solidFill = "0284C7"
s_mq.marker.graphicalProperties.line.solidFill = "0369A1"
s_mq.graphicalProperties.line.noFill = True
chart_mq.series.append(s_mq)

# 2.0 uM Warning Line
x_lim = Reference(ws, min_col=1, min_row=62, max_row=87)
y_lim = Reference(ws, min_col=4, min_row=62, max_row=87)
s_lim = Series(y_lim, x_lim, title="2.0 μM 警戒上限")
s_lim.marker.symbol = "none"
s_lim.graphicalProperties.line.solidFill = "DC2626"
s_lim.graphicalProperties.line.width = 20000
s_lim.graphicalProperties.line.prstDash = "dash"
chart_mq.series.append(s_lim)

if chart_mq.legend:
    chart_mq.legend.legendPos = "tr"

chart_mq.x_axis.title = "分析序列时序 (Sequence 1~26 / 柱 1~16)"
chart_mq.y_axis.title = "MQ 空白浓度 (μmol/L)"
chart_mq.y_axis.scaling.min = 0.0
chart_mq.y_axis.scaling.max = 3.5
chart_mq.x_axis.scaling.min = 0.5
chart_mq.x_axis.scaling.max = 26.5

ws.add_chart(chart_mq, "A37")

# -----------------------------------------------------------------------------
# Chart 2: DSW CRM Precision & Recovery Scatter Chart
# -----------------------------------------------------------------------------
chart_dsw = ScatterChart()
chart_dsw.title = "【图 2】全航段 26 序列 / 16 柱 DSW (CRM) 质控准确度与系统偏差总览散点图 (Live DSW Control)"
chart_dsw.style = 13
chart_dsw.width = 22.0
chart_dsw.height = 11.5

# DSW Data points (n=190)
x_dsw = Reference(ws, min_col=6, min_row=62, max_row=251)
y_dsw = Reference(ws, min_col=8, min_row=62, max_row=251)
s_dsw = Series(y_dsw, x_dsw, title="DSW (CRM) 测定点 (n=190)")
s_dsw.marker = Marker('square')
s_dsw.marker.size = 5
s_dsw.marker.graphicalProperties.solidFill = "10B981"
s_dsw.marker.graphicalProperties.line.solidFill = "047857"
s_dsw.graphicalProperties.line.noFill = True
chart_dsw.series.append(s_dsw)

# True Value Line (44.0 uM)
x_true = Reference(ws, min_col=6, min_row=62, max_row=87)
y_true = Reference(ws, min_col=9, min_row=62, max_row=87)
s_true = Series(y_true, x_true, title="44.0 μM 标称真值")
s_true.marker.symbol = "none"
s_true.graphicalProperties.line.solidFill = "1E293B"
s_true.graphicalProperties.line.width = 22000
s_true.graphicalProperties.line.prstDash = "solid"
chart_dsw.series.append(s_true)

# +5% Control Limit (46.2 uM)
y_p5 = Reference(ws, min_col=10, min_row=62, max_row=87)
s_p5 = Series(y_p5, x_true, title="+5% 控制限 (46.2 μM)")
s_p5.marker.symbol = "none"
s_p5.graphicalProperties.line.solidFill = "EA580C"
s_p5.graphicalProperties.line.width = 18000
s_p5.graphicalProperties.line.prstDash = "dash"
chart_dsw.series.append(s_p5)

# -5% Control Limit (41.8 uM)
y_m5 = Reference(ws, min_col=11, min_row=62, max_row=87)
s_m5 = Series(y_m5, x_true, title="-5% 控制限 (41.8 μM)")
s_m5.marker.symbol = "none"
s_m5.graphicalProperties.line.solidFill = "EA580C"
s_m5.graphicalProperties.line.width = 18000
s_m5.graphicalProperties.line.prstDash = "dash"
chart_dsw.series.append(s_m5)

if chart_dsw.legend:
    chart_dsw.legend.legendPos = "tr"

chart_dsw.x_axis.title = "分析序列时序 (Sequence 1~26 / 柱 1~16)"
chart_dsw.y_axis.title = "DSW 浓度 (μmol/L)"
chart_dsw.y_axis.scaling.min = 35.0
chart_dsw.y_axis.scaling.max = 55.0
chart_dsw.x_axis.scaling.min = 0.5
chart_dsw.x_axis.scaling.max = 26.5

ws.add_chart(chart_dsw, "G37")

# Save workbook
wb.save(target_file)
print("Successfully upgraded charts and saved target workbook!")
