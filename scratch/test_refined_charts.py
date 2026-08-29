import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

p = r'F:\印度洋测样\ODV\202608\20260827\Ocean_DOC_MultiColumn_QC_Report_GEOMAR_Validated_v2_Processed_latest.xlsx'
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb['Executive_Dashboard']
ws._charts.clear()

# -----------------------------------------------------------------------------
# Chart 1: MQ Baseline
# -----------------------------------------------------------------------------
chart_mq = ScatterChart()
chart_mq.title = "【图 1】全航段 26 序列 / 16 柱 MQ 空白基线总览散点图 (Live MQ Baseline)"
chart_mq.width = 23.5
chart_mq.height = 12.0

# 1. MQ Points (n=513)
x_mq = Reference(ws, min_col=1, min_row=62, max_row=574)
y_mq = Reference(ws, min_col=3, min_row=62, max_row=574)
s_mq = Series(y_mq, x_mq, title="MQ 空白测定点 (n=513)")
s_mq.marker = Marker('circle')
s_mq.marker.size = 5
s_mq.marker.graphicalProperties.solidFill = "0284C7"
s_mq.marker.graphicalProperties.line.solidFill = "0369A1"
s_mq.graphicalProperties.line.noFill = True
chart_mq.series.append(s_mq)

# 2. 2.0 uM Warning Line - MUST SPAN FULL RANGE (rows 62 to 574)
y_lim = Reference(ws, min_col=4, min_row=62, max_row=574)
s_lim = Series(y_lim, x_mq, title="2.0 μM 警戒上限")
s_lim.marker.symbol = "none"
s_lim.graphicalProperties.line.solidFill = "DC2626"
s_lim.graphicalProperties.line.width = 22000  # ~1.75 pt
s_lim.graphicalProperties.line.prstDash = "dash"
chart_mq.series.append(s_lim)

if chart_mq.legend:
    chart_mq.legend.legendPos = "tr"
    chart_mq.legend.overlay = False

chart_mq.x_axis.title = "分析序列时序 (Sequence 1~26 / 柱 1~16)"
chart_mq.y_axis.title = "MQ 空白浓度 (μmol/L)"
chart_mq.y_axis.scaling.min = 0.0
chart_mq.y_axis.scaling.max = 3.5
chart_mq.x_axis.scaling.min = 0.5
chart_mq.x_axis.scaling.max = 26.5

chart_mq.x_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="CBD5E1", prstDash="dash"))
chart_mq.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="CBD5E1", prstDash="dash"))

ws.add_chart(chart_mq, "A37")

# -----------------------------------------------------------------------------
# Chart 2: DSW CRM Precision
# -----------------------------------------------------------------------------
chart_dsw = ScatterChart()
chart_dsw.title = "【图 2】全航段 26 序列 / 16 柱 DSW (CRM) 质控准确度与系统偏差总览散点图 (Live DSW Control)"
chart_dsw.width = 23.5
chart_dsw.height = 12.0

# 1. DSW Points (n=190)
x_dsw = Reference(ws, min_col=6, min_row=62, max_row=251)
y_dsw = Reference(ws, min_col=8, min_row=62, max_row=251)
s_dsw = Series(y_dsw, x_dsw, title="DSW (CRM) 测定点 (n=190)")
s_dsw.marker = Marker('square')
s_dsw.marker.size = 5
s_dsw.marker.graphicalProperties.solidFill = "10B981"
s_dsw.marker.graphicalProperties.line.solidFill = "047857"
s_dsw.graphicalProperties.line.noFill = True
chart_dsw.series.append(s_dsw)

# 2. True Value Line (44.0 uM) - FULL RANGE (rows 62 to 251)
y_true = Reference(ws, min_col=9, min_row=62, max_row=251)
s_true = Series(y_true, x_dsw, title="44.0 μM 标称真值")
s_true.marker.symbol = "none"
s_true.graphicalProperties.line.solidFill = "1E293B"
s_true.graphicalProperties.line.width = 24000  # ~2 pt
s_true.graphicalProperties.line.prstDash = "solid"
chart_dsw.series.append(s_true)

# 3. +5% Control Limit (46.2 uM) - FULL RANGE
y_p5 = Reference(ws, min_col=10, min_row=62, max_row=251)
s_p5 = Series(y_p5, x_dsw, title="+5% 控制限 (46.2 μM)")
s_p5.marker.symbol = "none"
s_p5.graphicalProperties.line.solidFill = "EA580C"
s_p5.graphicalProperties.line.width = 18000
s_p5.graphicalProperties.line.prstDash = "dash"
chart_dsw.series.append(s_p5)

# 4. -5% Control Limit (41.8 uM) - FULL RANGE
y_m5 = Reference(ws, min_col=11, min_row=62, max_row=251)
s_m5 = Series(y_m5, x_dsw, title="-5% 控制限 (41.8 μM)")
s_m5.marker.symbol = "none"
s_m5.graphicalProperties.line.solidFill = "EA580C"
s_m5.graphicalProperties.line.width = 18000
s_m5.graphicalProperties.line.prstDash = "dash"
chart_dsw.series.append(s_m5)

if chart_dsw.legend:
    chart_dsw.legend.legendPos = "tr"
    chart_dsw.legend.overlay = False

chart_dsw.x_axis.title = "分析序列时序 (Sequence 1~26 / 柱 1~16)"
chart_dsw.y_axis.title = "DSW 浓度 (μmol/L)"
chart_dsw.y_axis.scaling.min = 35.0
chart_dsw.y_axis.scaling.max = 55.0
chart_dsw.x_axis.scaling.min = 0.5
chart_dsw.x_axis.scaling.max = 26.5

chart_dsw.x_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="CBD5E1", prstDash="dash"))
chart_dsw.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="CBD5E1", prstDash="dash"))

ws.add_chart(chart_dsw, "G37")

wb.save("scratch/test_refined_charts.xlsx")
print("Saved scratch/test_refined_charts.xlsx successfully!")
